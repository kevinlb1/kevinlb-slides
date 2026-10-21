"""Turn a slide's click-by-click animation into one slide per step.

LibreOffice does not play animations when it converts a deck to PDF: every
shape is drawn where it ends up, so anything that appears or disappears during
a slide is on the page at once, and a shape that appears over another hides it
for good. This module reads a slide's main animation sequence (the effects that
run on mouse clicks, with their "with previous" and "after previous"
companions) and replaces the slide with one copy per audience view: the state
before the first click and the state after each click. A copy omits the shapes
that have not yet appeared or have already gone, and draws the paragraphs of a
bulleted build that are still to come in fully transparent ink, so the visible
lines keep the positions and shrink-to-fit size they have on the finished
slide.

Only entrance and exit effects change what is on the page; emphasis, motion
paths, media and triggered (click-on-a-shape) sequences are ignored. Whether a
slide needs this at all is decided by `reason_for_steps`: any exit effect, or
an appearing shape that covers a shape already showing. The deck author can
force or forbid it with a tag in the slide's speaker notes; see
stamp_activity_qr.py.
"""

from __future__ import annotations

import copy
from dataclasses import dataclass
import re

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn

COVER_FRACTION = 0.25      # of the smaller shape's area
BACKDROP_FRACTION = 0.9    # of the slide: a shape this big is a backdrop, never "covered"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
FILL_TAGS = {qn(t) for t in ("a:noFill", "a:solidFill", "a:gradFill",
                             "a:blipFill", "a:pattFill", "a:grpFill")}
BULLET_TAGS = {qn(t) for t in ("a:buClrTx", "a:buClr", "a:buSzTx", "a:buSzPct", "a:buSzPts",
                               "a:buFontTx", "a:buFont", "a:buNone", "a:buAutoNum",
                               "a:buChar", "a:buBlip")}
AFTER_BULLET = {qn(t) for t in ("a:tabLst", "a:defRPr", "a:extLst")}
SLIDE_PARTNAME = re.compile(r"^/ppt/slides/slide(\d+)[.]xml$")
INVISIBLE_INK = ('<a:solidFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                 '<a:srgbClr val="000000"><a:alpha val="0"/></a:srgbClr></a:solidFill>')


@dataclass(frozen=True)
class Effect:
    kind: str                       # "entr" or "exit"
    shape_id: str
    paragraphs: frozenset | None    # None: the whole shape


@dataclass
class Group:
    click: bool
    effects: list


@dataclass(frozen=True)
class View:
    """What is missing from the slide at one moment."""
    hidden_shapes: frozenset
    hidden_paragraphs: tuple        # sorted ((shape_id, frozenset(indices)), ...)


def main_sequence(sld) -> list[Group]:
    """The click groups of the slide's main animation sequence, in order."""
    timing = sld.find(qn("p:timing"))
    if timing is None:
        return []
    for seq in timing.iter(qn("p:seq")):
        ctn = seq.find(qn("p:cTn"))
        if ctn is not None and ctn.get("nodeType") == "mainSeq":
            break
    else:
        return []
    groups = []
    children = ctn.find(qn("p:childTnLst"))
    for par in (children if children is not None else []):
        gctn = par.find(qn("p:cTn"))
        conds = gctn.findall(qn("p:stCondLst") + "/" + qn("p:cond")) if gctn is not None else []
        # A group that also starts when the slide begins ("after previous" on
        # the first effect) is not a click of its own.
        click = not any(cond.get("evt") == "onBegin" for cond in conds)
        effects = []
        for ctn in par.iter(qn("p:cTn")):
            kind = ctn.get("presetClass")
            if kind not in ("entr", "exit"):
                continue
            target = next(ctn.iter(qn("p:spTgt")), None)
            if target is None:
                continue
            rng = target.find(qn("p:txEl") + "/" + qn("p:pRg"))
            paragraphs = None if rng is None else frozenset(
                range(int(rng.get("st")), int(rng.get("end")) + 1))
            effects.append(Effect(kind, target.get("spid"), paragraphs))
        groups.append(Group(click, effects))
    return groups


def _initially_hidden(groups):
    """Anything with an entrance effect starts out unseen."""
    shapes, paragraphs = set(), {}
    for group in groups:
        for effect in group.effects:
            if effect.kind != "entr":
                continue
            if effect.paragraphs is None:
                shapes.add(effect.shape_id)
            else:
                paragraphs.setdefault(effect.shape_id, set()).update(effect.paragraphs)
    return shapes, paragraphs


def _apply(effect, shapes, paragraphs):
    if effect.paragraphs is None:
        (shapes.discard if effect.kind == "entr" else shapes.add)(effect.shape_id)
        return
    pending = paragraphs.setdefault(effect.shape_id, set())
    if effect.kind == "entr":
        pending -= effect.paragraphs
    else:
        pending |= effect.paragraphs


def views(groups) -> list[View]:
    """The distinct states the audience sees: before the first click and
    after each click. Effects that start with the slide fold into the view
    before them; clicks that change nothing visible add no view."""
    shapes, paragraphs = _initially_hidden(groups)

    def snapshot():
        return View(frozenset(shapes),
                    tuple(sorted((sid, frozenset(idx)) for sid, idx in paragraphs.items() if idx)))

    pages = [snapshot()]
    for group in groups:
        for effect in group.effects:
            _apply(effect, shapes, paragraphs)
        if group.click:
            pages.append(snapshot())
        else:
            pages[-1] = snapshot()
    distinct = [pages[0]]
    for view in pages[1:]:
        if view != distinct[-1]:
            distinct.append(view)
    return distinct


def geometry(slide):
    """Bounding box (left, top, width, height) and name of every shape by id.
    Shapes inside a group take the group's box, which is what matters here."""
    boxes, names = {}, {}

    def visit(shapes, inherited=None):
        for shape in shapes:
            box = inherited
            if box is None:
                try:
                    box = (shape.left, shape.top, shape.width, shape.height)
                except Exception:  # noqa: BLE001 - odd shapes without a frame
                    box = None
                if box is not None and any(v is None for v in box):
                    box = None
            sid = str(shape.shape_id)
            names[sid] = shape.name
            if box is not None:
                boxes[sid] = box
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                visit(shape.shapes, box)

    visit(slide.shapes)
    return boxes, names


def covers(a, b, slide_area) -> bool:
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    if min(aw, ah, bw, bh) <= 0 or bw * bh >= BACKDROP_FRACTION * slide_area:
        return False
    ix = max(0, min(ax + aw, bx + bw) - max(ax, bx))
    iy = max(0, min(ay + ah, by + bh) - max(ay, by))
    return ix * iy >= COVER_FRACTION * min(aw * ah, bw * bh)


def reason_for_steps(prs, slide) -> str | None:
    """Why the slide cannot be shown as a single page, or None if it can:
    something disappears, or something appears on top of what is showing."""
    groups = main_sequence(slide._element)
    if len(views(groups)) < 2:
        return None
    boxes, names = geometry(slide)
    slide_area = prs.slide_width * prs.slide_height
    hidden, paragraphs = _initially_hidden(groups)
    for group in groups:
        showing = [sid for sid in boxes if sid not in hidden]
        for effect in group.effects:
            label = names.get(effect.shape_id, f"shape {effect.shape_id}")
            if effect.kind == "exit":
                return f"{label} disappears"
            box = boxes.get(effect.shape_id)
            if box is None:
                continue
            for other in showing:
                if other != effect.shape_id and covers(box, boxes[other], slide_area):
                    return f"{label} covers {names[other]}"
        for effect in group.effects:
            _apply(effect, hidden, paragraphs)
    return None


def literalise_slide_numbers(prs) -> None:
    """Replace slide-number fields with the number each slide has in the deck.

    LibreOffice fills the field with the page's position, so once one slide
    becomes several pages every later page would be renumbered. Fixed text
    keeps the numbering of the deck, which is what the lecturer refers to."""
    for number, slide in enumerate(prs.slides, 1):
        for field in list(slide._element.iter(qn("a:fld"))):
            if field.get("type") != "slidenum":
                continue
            field.tag = qn("a:r")
            for attribute in ("id", "type"):
                field.attrib.pop(attribute, None)
            ppr = field.find(qn("a:pPr"))
            if ppr is not None:
                field.remove(ppr)
            text = field.find(qn("a:t"))
            if text is None:
                text = etree.SubElement(field, qn("a:t"))
            text.text = str(number)


def free_slide_partname(prs) -> PackURI:
    """python-pptx names a new slide part after the slide count, which falls
    once an original is dropped, so a later copy would reuse a name in the
    package and the saved file would carry two parts under it."""
    used = [int(m.group(1)) for part in prs.part.package.iter_parts()
            if (m := SLIDE_PARTNAME.match(str(part.partname)))]
    return PackURI(f"/ppt/slides/slide{max(used, default=0) + 1}.xml")


def flatten_ole_objects(cSld) -> int:
    """Swap embedded OLE objects for the picture PowerPoint keeps as their
    fallback. An OLE object hangs off a legacy VML drawing part that must
    belong to one slide, so a copy cannot share it; the fallback picture is
    what LibreOffice draws for the object in any case."""
    swapped = 0
    for frame in list(cSld.iter(qn("p:graphicFrame"))):
        if next(frame.iter(qn("p:oleObj")), None) is None:
            continue
        fallback = next(frame.iter(MC_FALLBACK), None)
        picture = None if fallback is None else next(fallback.iter(qn("p:pic")), None)
        if picture is None:
            continue
        picture = copy.deepcopy(picture)
        # The fallback's own id and frame are placeholders; take the object's.
        properties = frame.find(qn("p:nvGraphicFramePr") + "/" + qn("p:cNvPr"))
        target = picture.find(qn("p:nvPicPr") + "/" + qn("p:cNvPr"))
        if properties is not None and target is not None:
            for attribute in ("id", "name"):
                if properties.get(attribute) is not None:
                    target.set(attribute, properties.get(attribute))
        frame_xfrm = frame.find(qn("p:xfrm"))
        spPr = picture.find(qn("p:spPr"))
        if frame_xfrm is not None and spPr is not None and spPr.find(qn("a:xfrm")) is None:
            xfrm = etree.Element(qn("a:xfrm"))
            for child in frame_xfrm:
                xfrm.append(copy.deepcopy(child))
            spPr.insert(0, xfrm)
        frame.getparent().replace(frame, picture)
        swapped += 1
    return swapped


def clone_slide(prs, slide):
    """A copy of the slide, appended to the deck, sharing its layout and media.

    Only the relationships the copied XML still refers to come along, plus the
    legacy drawing of any OLE object that has no picture to fall back on;
    notes and comments stay with the original."""
    clone = prs.slides.add_slide(slide.slide_layout)
    clone.part.partname = free_slide_partname(prs)
    source, target = slide._element, clone._element
    for tag in ("p:cSld", "p:clrMapOvr", "p:transition", "p:timing", "p:extLst"):
        for element in target.findall(qn(tag)):
            target.remove(element)
    cSld = copy.deepcopy(source.find(qn("p:cSld")))
    flatten_ole_objects(cSld)
    target.insert(0, cSld)
    colour_map = source.find(qn("p:clrMapOvr"))
    if colour_map is not None:
        target.insert(1, copy.deepcopy(colour_map))

    referenced = {value for element in target.iter()
                  for name, value in element.attrib.items() if name.startswith("{" + RNS + "}")}
    keep_legacy = next(target.iter(qn("p:oleObj")), None) is not None
    mapping = {}
    for rel in slide.part.rels.values():
        if rel.rId not in referenced and not (keep_legacy and rel.reltype == RT.VML_DRAWING):
            continue
        if rel.is_external:
            mapping[rel.rId] = clone.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            mapping[rel.rId] = clone.part.relate_to(rel.target_part, rel.reltype)
    for element in target.iter():
        for name, value in element.attrib.items():
            if name.startswith("{" + RNS + "}") and value in mapping:
                element.set(name, mapping[value])
    return clone


def hide_paragraph(paragraph) -> None:
    """Keep the paragraph's space but draw nothing: transparent ink, no bullet."""
    for run in paragraph.findall(qn("a:r")) + paragraph.findall(qn("a:fld")):
        properties = run.find(qn("a:rPr"))
        if properties is None:
            properties = etree.Element(qn("a:rPr"))
            run.insert(0, properties)
        for child in list(properties):
            if child.tag in FILL_TAGS:
                properties.remove(child)
        outline = properties.find(qn("a:ln"))
        position = list(properties).index(outline) + 1 if outline is not None else 0
        properties.insert(position, etree.fromstring(INVISIBLE_INK))
    properties = paragraph.find(qn("a:pPr"))
    if properties is None:
        properties = etree.Element(qn("a:pPr"))
        paragraph.insert(0, properties)
    for child in list(properties):
        if child.tag in BULLET_TAGS:
            properties.remove(child)
    position = next((i for i, child in enumerate(properties) if child.tag in AFTER_BULLET),
                    len(properties))
    properties.insert(position, etree.Element(qn("a:buNone")))


def show_picture_fallbacks(prs) -> int:
    """Make LibreOffice draw the pictures PowerPoint keeps for text it cannot read.

    A text box holding an equation is saved twice: the text with its math,
    which LibreOffice stops drawing at the first equation, and a fallback
    copy of the box filled with a picture of the text, which LibreOffice
    uses. When that fallback is a placeholder with no text in it, LibreOffice
    takes it for an unfilled "click to add text" box and exports nothing,
    picture included, leaving the slide blank. One run of invisible text,
    without a bullet, makes it a filled box. Returns how many were changed."""
    changed = 0
    for slide in prs.slides:
        for fallback in slide._element.iter(MC_FALLBACK):
            for shape in fallback.iter(qn("p:sp")):
                spPr, body = shape.find(qn("p:spPr")), shape.find(qn("p:txBody"))
                if spPr is None or body is None or spPr.find(qn("a:blipFill")) is None:
                    continue
                paragraphs = body.findall(qn("a:p")) or [etree.SubElement(body, qn("a:p"))]
                if next(paragraphs[0].iter(qn("a:r")), None) is None:
                    run = etree.Element(qn("a:r"))
                    etree.SubElement(run, qn("a:t")).text = " "
                    end = paragraphs[0].find(qn("a:endParaRPr"))
                    if end is None:
                        paragraphs[0].append(run)
                    else:
                        end.addprevious(run)
                for paragraph in paragraphs:
                    hide_paragraph(paragraph)
                changed += 1
    return changed


def apply_view(sld, view: View) -> None:
    for tag in ("p:transition", "p:timing"):
        for element in sld.findall(qn(tag)):
            sld.remove(element)
    by_id = {nv.get("id"): nv for nv in sld.iter(qn("p:cNvPr"))}
    for sid in view.hidden_shapes:
        nv = by_id.get(sid)
        if nv is not None:
            shape = nv.getparent().getparent()
            shape.getparent().remove(shape)
    for sid, indices in view.hidden_paragraphs:
        nv = by_id.get(sid)
        if nv is None:
            continue
        body = nv.getparent().getparent().find(qn("p:txBody"))
        if body is None:
            continue
        paragraphs = body.findall(qn("a:p"))
        for index in indices:
            if index < len(paragraphs):
                hide_paragraph(paragraphs[index])


def expand(prs, slide) -> int:
    """Replace the slide with one copy per view; returns how many."""
    pages = views(main_sequence(slide._element))
    if len(pages) < 2:
        return 1
    id_list = prs.slides._sldIdLst
    position = prs.slides.index(slide)
    for offset, view in enumerate(pages):
        clone = clone_slide(prs, slide)
        apply_view(clone._element, view)
        entry = id_list[-1]
        id_list.remove(entry)
        id_list.insert(position + 1 + offset, entry)
    original = id_list[position]
    prs.part.drop_rel(original.rId)
    id_list.remove(original)
    return len(pages)
